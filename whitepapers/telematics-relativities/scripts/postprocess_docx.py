"""Apply the narrative-proposal style and editorial cover to the Pandoc DOCX."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = RGBColor(46, 74, 98)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(100, 108, 117)
LIGHT_FILL = "F4F6F9"
FONT = "Calibri"


def set_run_font(run, name=FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def get_style(doc, name):
    """Resolve Pandoc styles even when python-docx's name index is unavailable."""
    for style in doc.styles:
        if style.name == name:
            return style
    raise KeyError(f"Style not found: {name}")


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])
    set_run_font(run, size=9, color=MUTED)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_styles(doc):
    normal = get_style(doc, "Normal")
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    settings = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in settings.items():
        style = get_style(doc, name)
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("TOC 1", "TOC 2"):
        try:
            style = get_style(doc, name)
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(9.2)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = True

        header = section.header
        p = header.paragraphs[0]
        p.text = "Bayesian Credibility and Telematics Relativities"
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            set_run_font(run, size=9, color=MUTED)

        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.clear()
        add_page_field(footer_p)

        section.first_page_header.paragraphs[0].text = ""
        section.first_page_footer.paragraphs[0].text = ""


def page_break_paragraph():
    """Return a native page-break paragraph for insertion into document XML."""
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run.append(br)
    paragraph.append(run)
    return paragraph


def move_toc_after_cover(doc):
    """Place Pandoc's TOC after the editorial cover instead of before it."""
    body = doc._element.body
    toc = next((node for node in list(body) if node.tag == qn("w:sdt")), None)
    if toc is None:
        raise RuntimeError("Pandoc table of contents was not found in the DOCX body.")

    # The white paper's numbered sections, appendix titles and reference title
    # all use Heading 2 beneath the editorial Title style. Appendix subsections
    # use Heading 3 below, so a level-2 field yields a concise document map
    # without listing every calculation note.
    for instr in toc.iter(qn("w:instrText")):
        if instr.text and 'TOC' in instr.text:
            # Do not retain Pandoc's ``\\u`` switch here. It makes Word add
            # every paragraph carrying an outline level, including appendix
            # subheadings whose named style falls outside the requested range.
            # Restricting the field to the document-level heading keeps the map
            # on one page and prevents a trailing, visually blank continuation.
            instr.text = 'TOC \\o "2-2" \\h \\z'

    cover_end = None
    for paragraph in doc.paragraphs:
        if "Date: August 2026" in paragraph.text:
            cover_end = paragraph._p
            break
    if cover_end is None:
        raise RuntimeError("Cover metadata paragraph was not found.")

    body.remove(toc)
    cover_end.addnext(toc)
    cover_end.addnext(page_break_paragraph())


def style_cover_and_paragraphs(doc):
    title = "Bayesian Credibility and Exposure-Normalised Telematics Relativities"
    subtitle = "A unified actuarial architecture for pricing, reserving, capital and risk transfer in gig-economy motor insurance"
    reference_mode = False
    appendix_mode = False

    for p in doc.paragraphs:
        text = p.text.strip()
        if text == title:
            p.style = get_style(doc, "Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(118)
            p.paragraph_format.space_after = Pt(12)
            for run in p.runs:
                set_run_font(run, size=28, color=BLUE, bold=True)
        elif text == subtitle:
            p.style = get_style(doc, "Subtitle")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(28)
            for run in p.runs:
                set_run_font(run, size=14, color=DARK_BLUE, italic=True)
        elif text.startswith("Author:") or text.startswith("Contact:") or text.startswith("Status:") or text.startswith("Date:"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                set_run_font(run, size=10.5, color=MUTED)
        elif text == "Abstract":
            p.paragraph_format.page_break_before = True
        elif text.startswith("Appendix "):
            p.style = get_style(doc, "Heading 2")
            p.paragraph_format.page_break_before = True
            appendix_mode = True
            reference_mode = False
        elif text == "References":
            p.style = get_style(doc, "Heading 2")
            p.paragraph_format.page_break_before = True
            appendix_mode = False
            reference_mode = True
        elif (
            appendix_mode
            and len(text) > 2
            and text[0] in "ABC"
            and text[1] == "."
            and text[2].isdigit()
        ):
            # These headings are semantically subordinate to Appendix A, B or
            # C and should not occupy the executive table of contents.
            p.style = get_style(doc, "Heading 3")
        elif text.startswith("Figure ") and any(run.bold for run in p.runs):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(10)
            for run in p.runs:
                set_run_font(run, size=9.5, color=MUTED, italic=True)
        elif reference_mode and text.startswith("["):
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            # A compact but still clearly separated bibliography prevents a
            # single final entry from being stranded on an otherwise empty
            # page while preserving one reference per paragraph.
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.03
            for run in p.runs:
                set_run_font(run, size=9)

        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)


def style_tables(doc):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        rows = table.rows
        if not rows:
            continue
        columns = len(rows[0].cells)
        width = Inches(6.5 / columns)
        for r_idx, row in enumerate(rows):
            for cell in row.cells:
                cell.width = width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                if r_idx == 0:
                    shade_cell(cell, LIGHT_FILL)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_run_font(run, size=9.2, bold=(True if r_idx == 0 else None))


def cap_images(doc):
    max_width = Inches(6.3)
    max_height = Inches(7.2)
    for shape in doc.inline_shapes:
        ratio = max_width / shape.width
        shape.width = max_width
        shape.height = int(shape.height * ratio)
        if shape.height > max_height:
            ratio = max_height / shape.height
            shape.height = max_height
            shape.width = int(shape.width * ratio)


def set_metadata(doc):
    cp = doc.core_properties
    cp.title = "Bayesian Credibility and Exposure-Normalised Telematics Relativities"
    cp.subject = "Actuarial telematics pricing, reserving, economic capital and risk transfer"
    cp.author = "Neville Maloba"
    cp.keywords = "actuarial; telematics; credibility; Bayesian; reserving; economic capital; reinsurance"
    cp.comments = "Generated from the canonical paper.md source."


def main() -> int:
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    doc = Document(source)
    configure_styles(doc)
    configure_sections(doc)
    style_cover_and_paragraphs(doc)
    move_toc_after_cover(doc)
    style_tables(doc)
    cap_images(doc)
    set_metadata(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Styled {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
